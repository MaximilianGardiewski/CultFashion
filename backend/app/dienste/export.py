"""Inventurergebnis als Excel - das Arbeitsergebnis fuer advarics.

Vier Blaetter: Differenzen (das, was bearbeitet wird), Vollstaendig,
Unbekannte Codes und das Zaehl-Log. Das Log ist nicht Beiwerk: bei einer
strittigen Differenz ist es die einzige Stelle, an der steht, wer wann was
gezaehlt hat.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.tabellen import Inventur, ScanEvent, Sollposition, Zaehlbereich
from app.dienste import auswertung as auswertung_modul

SCHRIFT = "Arial"
KOPF_FUELLUNG = PatternFill("solid", fgColor="1F3864")
KOPF_SCHRIFT = Font(name=SCHRIFT, size=10, bold=True, color="FFFFFF")
TEXT = Font(name=SCHRIFT, size=10)
TITEL = Font(name=SCHRIFT, size=12, bold=True)
ROT = Font(name=SCHRIFT, size=10, color="B03030")
GRUEN = Font(name=SCHRIFT, size=10, color="1E7B34")


def _kopf(ws, spalten: list[tuple[str, int]], zeile: int = 1) -> None:
    for i, (name, breite) in enumerate(spalten, start=1):
        z = ws.cell(row=zeile, column=i, value=name)
        z.font = KOPF_SCHRIFT
        z.fill = KOPF_FUELLUNG
        z.alignment = Alignment(horizontal="left", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = breite
    ws.row_dimensions[zeile].height = 20
    ws.freeze_panes = ws.cell(row=zeile + 1, column=1)


def _schreibe_position(ws, r: int, d) -> None:
    p = d.position
    werte = [p.ean or "", p.artikelnummer, p.artikelname, p.marke,
             p.farbnummer, p.farbe, p.groesse, float(p.vk_preis or 0),
             d.buchbestand, d.gezaehlt, d.differenz, d.wert]
    for c, wert in enumerate(werte, start=1):
        z = ws.cell(row=r, column=c, value=wert)
        z.font = TEXT
        if c in (1, 2, 5, 7):
            z.number_format = "@"
        elif c == 8:
            z.number_format = '#,##0.00 "€"'
        elif c == 12:
            z.number_format = '#,##0.00 "€";[Red]-#,##0.00 "€"'
        if c in (11, 12) and d.differenz:
            z.font = ROT if d.differenz < 0 else GRUEN


SPALTEN_DIFF = [("EAN", 16), ("Artikelnummer", 15), ("Artikelname", 24), ("Marke", 13),
                ("Farbnummer", 11), ("Farbe", 18), ("Größe", 9), ("VK-Preis", 11),
                ("Buchbestand", 12), ("Gezählt", 10), ("Differenz", 11), ("Wert", 12)]


def baue(sitzung: Session, inventur: Inventur) -> BytesIO:
    a = auswertung_modul.erstelle(sitzung, inventur, nur_abweichungen=False)
    abweichungen = [d for d in a.differenzen if d.differenz != 0]

    wb = Workbook()

    # -- Differenzen --------------------------------------------------
    ws = wb.active
    ws.title = "Differenzen"
    ws["A1"] = f"Inventurdifferenzen – {inventur.bezeichnung}"
    ws["A1"].font = TITEL
    ws["A2"] = (f"Filiale {inventur.filial_nr} {inventur.filiale} · "
                f"{len(abweichungen)} Abweichungen · "
                f"Fehlmenge {a.kennzahlen.fehlmenge} · "
                f"Überbestand {a.kennzahlen.ueberbestand} · "
                f"Wert {a.kennzahlen.differenz_wert:.2f} €".replace(".", ","))
    ws["A2"].font = Font(name=SCHRIFT, size=9, italic=True, color="666666")
    _kopf(ws, SPALTEN_DIFF, zeile=4)
    for r, d in enumerate(abweichungen, start=5):
        _schreibe_position(ws, r, d)

    # -- Vollstaendig -------------------------------------------------
    ws = wb.create_sheet("Vollständig")
    _kopf(ws, SPALTEN_DIFF)
    for r, d in enumerate(a.differenzen, start=2):
        _schreibe_position(ws, r, d)

    # -- Unbekannte Codes ---------------------------------------------
    ws = wb.create_sheet("Unbekannte Codes")
    ws["A1"] = "Gescannt, aber nicht im Sollbestand"
    ws["A1"].font = TITEL
    ws["A2"] = ("Diese Teile lagen im Laden, advarics kennt sie unter dieser EAN nicht. "
                "Vor der Verbuchung klären.")
    ws["A2"].font = Font(name=SCHRIFT, size=9, italic=True, color="666666")
    _kopf(ws, [("EAN", 18), ("Roh-Code", 18), ("Anzahl", 10)], zeile=4)
    for r, u in enumerate(a.unbekannte, start=5):
        for c, wert in enumerate([u.ean or "", u.roh_code or "", u.menge], start=1):
            z = ws.cell(row=r, column=c, value=wert)
            z.font = TEXT
            if c <= 2:
                z.number_format = "@"

    # -- Zaehl-Log ----------------------------------------------------
    ws = wb.create_sheet("Zähl-Log")
    _kopf(ws, [("Zeit", 18), ("Bereich", 16), ("Erfasst von", 16), ("Art", 11),
               ("EAN", 16), ("Artikel", 34), ("Farbe", 16), ("Größe", 9),
               ("Menge", 8)])

    bereiche = {b.id: b.name for b in sitzung.scalars(
        select(Zaehlbereich).where(Zaehlbereich.inventur_id == inventur.id))}

    events = sitzung.execute(
        select(ScanEvent, Sollposition)
        .outerjoin(Sollposition, ScanEvent.sollposition_id == Sollposition.id)
        .where(ScanEvent.inventur_id == inventur.id)
        .order_by(ScanEvent.id)
    ).all()

    for r, (ev, pos) in enumerate(events, start=2):
        werte = [
            ev.erfasst_am.strftime("%d.%m.%Y %H:%M:%S") if ev.erfasst_am else "",
            bereiche.get(ev.zaehlbereich_id, ""),
            ev.erfasst_von,
            ev.erfassungsart,
            ev.ean or ev.roh_code or "",
            f"{pos.marke} {pos.artikelname}" if pos else "— unbekannt —",
            pos.farbe if pos else "",
            pos.groesse if pos else "",
            ev.menge,
        ]
        for c, wert in enumerate(werte, start=1):
            z = ws.cell(row=r, column=c, value=wert)
            z.font = TEXT
            if c in (5, 8):
                z.number_format = "@"

    puffer = BytesIO()
    wb.save(puffer)
    puffer.seek(0)
    return puffer
